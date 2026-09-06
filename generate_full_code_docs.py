#!/usr/bin/env python3
"""Generate comprehensive Word doc with full code for Experiments 2-8."""

import os
import subprocess
from pathlib import Path
from docx import Document
from docx.shared import Pt

BASE = Path("/home/mani-arch/Desktop/FSD lab/fullstack-web-lab")
OUT = BASE / "documents"
OUT.mkdir(exist_ok=True)

EXPS = [
    {
        "num": 2, "name": "Todo Application",
        "desc": "A full-stack Todo Management application allowing users to create, read, update, and delete tasks with priority levels, categories, due dates, and completion status.",
        "frontend_dir": "frontend/lib/features/experiment2_todo",
        "backend_files": [
            "backend/app/api/routes/todos.py",
            "backend/app/schemas/todo.py",
            "backend/app/services/todo_service.py",
            "backend/tests/test_todos.py",
        ],
        "test_file": "backend/tests/test_todos.py",
    },
    {
        "num": 3, "name": "Micro Blogging Platform",
        "desc": "A social micro-blogging platform where users can create short posts, follow other users, like posts, edit their profile, and view a personalized feed.",
        "frontend_dir": "frontend/lib/features/experiment3_blog",
        "backend_files": [
            "backend/app/api/routes/posts.py",
            "backend/app/api/routes/follows.py",
            "backend/app/api/routes/users.py",
            "backend/app/schemas/post.py",
            "backend/app/services/post_service.py",
            "backend/tests/test_posts.py",
        ],
        "test_file": "backend/tests/test_posts.py",
    },
    {
        "num": 4, "name": "Food Delivery System",
        "desc": "A food delivery platform with restaurant browsing, menu viewing, cart management, order placement, and order tracking.",
        "frontend_dir": "frontend/lib/features/experiment4_food",
        "backend_files": [
            "backend/app/api/routes/restaurants.py",
            "backend/app/api/routes/food.py",
            "backend/app/api/routes/cart.py",
            "backend/app/api/routes/orders.py",
            "backend/app/schemas/food.py",
            "backend/app/services/food_service.py",
            "backend/tests/test_food.py",
        ],
        "test_file": "backend/tests/test_food.py",
    },
    {
        "num": 5, "name": "Classifieds Marketplace",
        "desc": "An online classifieds marketplace for buying and selling items with listing creation, categories, favorites, messaging, and admin moderation.",
        "frontend_dir": "frontend/lib/features/experiment5_classifieds",
        "backend_files": [
            "backend/app/api/routes/listings.py",
            "backend/app/api/routes/favorites.py",
            "backend/app/schemas/classified.py",
            "backend/app/services/classified_service.py",
            "backend/tests/test_classifieds.py",
        ],
        "test_file": "backend/tests/test_classifieds.py",
    },
    {
        "num": 6, "name": "Leave Management System",
        "desc": "An employee leave management system with leave request submission, approval workflow, balance tracking, and admin dashboards.",
        "frontend_dir": "frontend/lib/features/experiment6_leave",
        "backend_files": [
            "backend/app/api/routes/leaves.py",
            "backend/app/api/routes/admin_leaves.py",
            "backend/app/schemas/leave.py",
            "backend/app/services/leave_service.py",
            "backend/tests/test_leaves.py",
        ],
        "test_file": "backend/tests/test_leaves.py",
    },
    {
        "num": 7, "name": "Project Management Tool",
        "desc": "A project management tool with project creation, task assignment, status tracking, and progress monitoring.",
        "frontend_dir": "frontend/lib/features/experiment7_project",
        "backend_files": [
            "backend/app/api/routes/projects.py",
            "backend/app/api/routes/tasks.py",
            "backend/app/schemas/project.py",
            "backend/app/services/project_service.py",
            "backend/tests/test_tasks.py",
        ],
        "test_file": "backend/tests/test_tasks.py",
    },
    {
        "num": 8, "name": "Online Survey System",
        "desc": "An online survey system for creating, distributing, and analyzing surveys with question bank management and auto-scoring.",
        "frontend_dir": "frontend/lib/features/experiment8_survey",
        "backend_files": [
            "backend/app/api/routes/survey.py",
            "backend/app/api/routes/questions.py",
            "backend/app/schemas/survey.py",
            "backend/app/services/survey_service.py",
            "backend/tests/test_survey.py",
        ],
        "test_file": "backend/tests/test_survey.py",
    },
]


def read_file(path):
    full = BASE / path
    try:
        with open(full, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    except FileNotFoundError:
        return f"// File not found: {path}"


def add_code(doc, filename, code, font_size=9):
    p = doc.add_paragraph()
    r = p.add_run(filename)
    r.bold = True; r.font.size = Pt(11); r.font.name = "Cambria"
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(1)
    p2.paragraph_format.line_spacing = Pt(13)
    for line in code.split("\n"):
        r = p2.add_run(line + "\n")
        r.font.name = "Cambria"
        r.font.size = Pt(font_size)


doc = Document()
style = doc.styles["Normal"]
style.font.name = "Cambria"
style.font.size = Pt(10)

# Title page
for _ in range(5):
    doc.add_paragraph("")
t = doc.add_paragraph(); t.alignment = 1
r = t.add_run("Full Stack Web Lab"); r.bold = True; r.font.size = Pt(26); r.font.name = "Cambria"
s = doc.add_paragraph(); s.alignment = 1
r = s.add_run("Experiments 2 to 8 — Complete Source Code"); r.bold = True; r.font.size = Pt(16); r.font.name = "Cambria"
doc.add_paragraph("")
info = doc.add_paragraph(); info.alignment = 1
r = info.add_run("Flutter Web + FastAPI + MongoDB\nAll frontend and backend source code\nFont: Cambria 10pt"); r.font.size = Pt(12); r.font.name = "Cambria"
doc.add_page_break()

# ====== SHARED FILES (before experiments) ======
p = doc.add_paragraph()
r = p.add_run("SHARED FRONTEND FILES"); r.bold = True; r.font.size = Pt(14); r.font.name = "Cambria"
doc.add_paragraph("")

shared_frontend = [
    "frontend/lib/main.dart",
    "frontend/lib/app_router.dart",
    "frontend/lib/lab_shell.dart",
    "frontend/lib/widgets/auth_gate.dart",
    "frontend/lib/widgets/app_shell.dart",
    "frontend/lib/widgets/feature_navigator.dart",
    "frontend/lib/services/session.dart",
    "frontend/lib/core/network/api_client.dart",
    "frontend/lib/core/constants/app_constants.dart",
    "frontend/lib/core/theme/app_theme.dart",
    "frontend/lib/core/utils/validators.dart",
    "frontend/lib/core/utils/formatters.dart",
    "frontend/lib/core/utils/date_utils.dart",
    "frontend/lib/core/storage/token_storage.dart",
    "frontend/lib/core/storage/app_cache.dart",
    "frontend/lib/models/user.dart",
    "frontend/lib/widgets/loading_widget.dart",
    "frontend/lib/widgets/error_widget.dart",
    "frontend/lib/widgets/empty_state.dart",
    "frontend/lib/widgets/animated_card.dart",
    "frontend/lib/widgets/stat_card.dart",
    "frontend/lib/widgets/status_chip.dart",
    "frontend/lib/widgets/app_dialogs.dart",
    "frontend/lib/features/dashboard/dashboard_screen.dart",
    "frontend/web/index.html",
    "frontend/web/manifest.json",
    "frontend/pubspec.yaml",
    "frontend/serve.py",
]

for fpath in shared_frontend:
    code = read_file(fpath)
    if code and not code.startswith("// File not found"):
        add_code(doc, fpath, code, font_size=9)
        doc.add_paragraph("")

# Shared backend
p = doc.add_paragraph()
r = p.add_run("SHARED BACKEND FILES"); r.bold = True; r.font.size = Pt(14); r.font.name = "Cambria"
doc.add_paragraph("")

shared_backend = [
    "backend/app/main.py",
    "backend/app/core/config.py",
    "backend/app/core/database.py",
    "backend/app/core/security.py",
    "backend/app/middleware/rate_limit.py",
    "backend/app/api/deps.py",
    "backend/app/utils/errors.py",
    "backend/app/utils/helpers.py",
    "backend/requirements.txt",
]

for fpath in shared_backend:
    code = read_file(fpath)
    if code and not code.startswith("// File not found"):
        add_code(doc, fpath, code, font_size=9)
        doc.add_paragraph("")

# Backend models and repositories
p = doc.add_paragraph()
r = p.add_run("BACKEND MODELS & REPOSITORIES"); r.bold = True; r.font.size = Pt(14); r.font.name = "Cambria"
doc.add_paragraph("")

backend_shared_py = [
    "backend/app/models/__init__.py",
    "backend/app/repositories/base.py",
    "backend/app/repositories/posts_repo.py",
    "backend/app/repositories/users_repo.py",
    "backend/app/repositories/survey_repo.py",
    "backend/app/repositories/classified_repo.py",
    "backend/app/repositories/food_repo.py",
    "backend/app/repositories/leaves_repo.py",
    "backend/app/repositories/projects_repo.py",
    "backend/requirements.txt",
    "backend/pytest.ini",
    "backend/scripts/seed_database.py",
    "backend/tests/conftest.py",
]

for fpath in backend_shared_py:
    code = read_file(fpath)
    if code and not code.startswith("// File not found"):
        add_code(doc, fpath, code, font_size=9)
        doc.add_paragraph("")

# ====== EXPERIMENTS ======
for exp in EXPS:
    h = doc.add_paragraph()
    r = h.add_run(f"EX.NO: {exp['num']}    {exp['name']}")
    r.bold = True; r.font.size = Pt(14); r.font.name = "Cambria"
    p = doc.add_paragraph()
    r = p.add_run(exp["desc"])
    r.italic = True; r.font.size = Pt(10); r.font.name = "Cambria"
    doc.add_paragraph("")

    # Frontend
    p = doc.add_paragraph()
    r = p.add_run("FRONTEND CODE (Flutter/Dart)"); r.bold = True; r.font.size = Pt(12); r.font.name = "Cambria"
    doc.add_paragraph("")

    result = subprocess.run(
        ["find", str(BASE / exp["frontend_dir"]), "-name", "*.dart"],
        capture_output=True, text=True
    )
    dart_files = sorted(result.stdout.strip().split("\n"))

    for fpath in dart_files:
        rel = os.path.relpath(fpath, BASE)
        code = read_file(rel)
        add_code(doc, rel, code, font_size=9)
        doc.add_paragraph("")

    # Backend
    p = doc.add_paragraph()
    r = p.add_run("BACKEND CODE (FastAPI/Python)"); r.bold = True; r.font.size = Pt(12); r.font.name = "Cambria"
    doc.add_paragraph("")

    for bpath in exp["backend_files"]:
        code = read_file(bpath)
        add_code(doc, bpath, code, font_size=9)
        doc.add_paragraph("")

    # Test files
    test_file = exp.get("test_file", "")
    if test_file:
        p = doc.add_paragraph()
        r = p.add_run(f"TESTS ({test_file.split('/')[-1].replace('.py', '').upper()})"); r.bold = True; r.font.size = Pt(12); r.font.name = "Cambria"
        doc.add_paragraph("")
        code = read_file(test_file)
        add_code(doc, test_file, code, font_size=9)
        doc.add_paragraph("")

    # OUTPUT
    p = doc.add_paragraph()
    r = p.add_run("OUTPUT :"); r.bold = True; r.font.size = Pt(12); r.font.name = "Cambria"
    for _ in range(40):
        doc.add_paragraph("")

    # RESULT
    p = doc.add_paragraph()
    r = p.add_run("RESULT :"); r.bold = True; r.font.size = Pt(12); r.font.name = "Cambria"
    doc.add_paragraph("")

    doc.add_page_break()

filepath = OUT / "All_Experiments_2_to_8_Complete_Code.docx"
doc.save(str(filepath))
print(f"Generated: {filepath.name}")