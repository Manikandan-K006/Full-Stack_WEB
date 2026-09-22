#!/usr/bin/env python3
"""Reduced code version with clear headings for Experiments 2-8."""

import os, subprocess
from pathlib import Path
from docx import Document
from docx.shared import Pt

BASE = Path("/home/mani-arch/Desktop/FSD lab/fullstack-web-lab")
OUT = BASE / "documents"
OUT.mkdir(exist_ok=True)

EXPS = [
    (2, "TODO APPLICATION", "Manage the To-Do list of users where users can login and manage their To-Do items", "frontend/lib/features/experiment2_todo", ["backend/app/api/routes/todos.py","backend/app/schemas/todo.py","backend/app/services/todo_service.py"]),
    (3, "MICRO BLOGGING APPLICATION", "Allows people to post their content which can be viewed by people who follow them", "frontend/lib/features/experiment3_blog", ["backend/app/api/routes/posts.py","backend/app/schemas/post.py","backend/app/services/post_service.py"]),
    (4, "FOOD DELIVERY WEBSITE", "Users can order food from a particular restaurant listed in the website", "frontend/lib/features/experiment4_food", ["backend/app/api/routes/restaurants.py","backend/app/api/routes/food.py","backend/app/schemas/food.py","backend/app/services/food_service.py"]),
    (5, "CLASSIFIEDS WEB APPLICATION", "Buy and sell used products", "frontend/lib/features/experiment5_classifieds", ["backend/app/api/routes/listings.py","backend/app/schemas/classified.py","backend/app/services/classified_service.py"]),
    (6, "LEAVE MANAGEMENT SYSTEM", "Manage employee leave requests", "frontend/lib/features/experiment6_leave", ["backend/app/api/routes/leaves.py","backend/app/schemas/leave.py","backend/app/services/leave_service.py"]),
    (7, "PROJECT MANAGEMENT TOOL", "Tracking tasks and projects", "frontend/lib/features/experiment7_project", ["backend/app/api/routes/projects.py","backend/app/schemas/project.py","backend/app/services/project_service.py"]),
    (8, "ONLINE SURVEY SYSTEM", "Create and manage surveys", "frontend/lib/features/experiment8_survey", ["backend/app/api/routes/survey.py","backend/app/schemas/survey.py","backend/app/services/survey_service.py"]),
]

def read_truncated(path, max_lines=28):
    full = BASE / path
    try:
        with open(full, "r", encoding="utf-8", errors="replace") as f:
            lines = f.read().split("\n")
            if len(lines) > max_lines:
                return "\n".join(lines[:max_lines]) + f"\n... ({len(lines)-max_lines} more lines truncated)"
            return "\n".join(lines)
    except FileNotFoundError:
        return f"// File not found: {path}"

def add_code(doc, filename, code):
    p = doc.add_paragraph()
    r = p.add_run(filename); r.bold = True; r.font.size = Pt(11); r.font.name = "Cambria"
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(4); p2.paragraph_format.line_spacing = Pt(14)
    for line in code.split("\n"):
        r = p2.add_run(line + "\n"); r.font.name = "Cambria"; r.font.size = Pt(10)

doc = Document()
style = doc.styles["Normal"]; style.font.name = "Cambria"; style.font.size = Pt(10)

# Title page
for _ in range(4): doc.add_paragraph("")
t = doc.add_paragraph(); t.alignment = 1; r = t.add_run("Full Stack Web Lab"); r.bold = True; r.font.size = Pt(26); r.font.name = "Cambria"
s = doc.add_paragraph(); s.alignment = 1; r = s.add_run("Experiments 2 to 8 — Reduced Code with Headings"); r.bold = True; r.font.size = Pt(14); r.font.name = "Cambria"
doc.add_paragraph("")
info = doc.add_paragraph(); info.alignment = 1; r = info.add_run("Flutter Web + FastAPI + MongoDB\nFont: Cambria | Code truncated to ~28 lines per file"); r.font.size = Pt(11); r.font.name = "Cambria"
# Table of contents style headings list
doc.add_paragraph("")
p = doc.add_paragraph(); r = p.add_run("CONTENTS"); r.bold = True; r.font.size = Pt(13); r.font.name = "Cambria"
for num, name, _, _, _ in EXPS:
    p = doc.add_paragraph(f"  Experiment {num} : {name}", style="List Bullet")
    for run in p.runs: run.font.name = "Cambria"; run.font.size = Pt(11)
doc.add_page_break()

for num, name, desc, fdir, bfiles in EXPS:
    # Main experiment heading - prominent
    h = doc.add_paragraph(); h.alignment = 1
    r = h.add_run(f"EXPERIMENT {num}"); r.bold = True; r.font.size = Pt(18); r.font.name = "Cambria"
    h2 = doc.add_paragraph(); h2.alignment = 1
    r = h2.add_run(name); r.bold = True; r.font.size = Pt(16); r.font.name = "Cambria"
    p = doc.add_paragraph(); p.alignment = 1
    r = p.add_run(desc); r.italic = True; r.font.size = Pt(10); r.font.name = "Cambria"
    doc.add_paragraph("")
    p = doc.add_paragraph(); p.alignment = 2
    r = p.add_run("9665"); r.bold = True; r.font.size = Pt(11); r.font.name = "Cambria"
    doc.add_paragraph("")
    p = doc.add_paragraph(); r = p.add_run("PROGRAM :"); r.bold = True; r.font.size = Pt(12); r.font.name = "Cambria"
    doc.add_paragraph("")

    # Frontend - reduced
    p = doc.add_paragraph(); r = p.add_run("FRONTEND CODE"); r.bold = True; r.font.size = Pt(12); r.font.name = "Cambria"; r.underline = True
    doc.add_paragraph("")
    result = subprocess.run(["find", str(BASE / fdir), "-name", "*.dart"], capture_output=True, text=True)
    for fpath in sorted(result.stdout.strip().split("\n")):
        if not fpath: continue
        rel = os.path.relpath(fpath, BASE)
        add_code(doc, rel, read_truncated(rel, 28))
        doc.add_paragraph("")

    # Backend - reduced
    p = doc.add_paragraph(); r = p.add_run("BACKEND CODE"); r.bold = True; r.font.size = Pt(12); r.font.name = "Cambria"; r.underline = True
    doc.add_paragraph("")
    for bf in bfiles:
        add_code(doc, bf, read_truncated(bf, 30))
        doc.add_paragraph("")

    # Output
    p = doc.add_paragraph(); r = p.add_run("OUTPUT :"); r.bold = True; r.font.size = Pt(12); r.font.name = "Cambria"
    for _ in range(18): doc.add_paragraph("")
    p = doc.add_paragraph(); r = p.add_run("RESULT :"); r.bold = True; r.font.size = Pt(12); r.font.name = "Cambria"
    doc.add_paragraph("")
    doc.add_page_break()

out = OUT / "All_Experiments_2_to_8_Reduced_Code.docx"
doc.save(str(out))
print(f"Generated: {out.name}")
